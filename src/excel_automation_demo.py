from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


def format_workbook(input_file, output_file):
    """
    Demonstrates selected Excel automation capabilities
    using OpenPyXL.
    """

    workbook = load_workbook(input_file)
    worksheet = workbook.active

    # Format header row
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Freeze header row
    worksheet.freeze_panes = "A2"

    # Automatically size columns
    for column in worksheet.columns:
        max_length = max(
            len(str(cell.value or ""))
            for cell in column
        )

        column_letter = get_column_letter(column[0].column)
        worksheet.column_dimensions[column_letter].width = max_length + 2

    workbook.save(output_file)
